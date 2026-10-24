import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_report():
    doc = Document()

    # Page Margins (ГОСТ: левое 30 мм, правое 15 мм, верхнее 20 мм, нижнее 20 мм)
    sections = doc.sections
    for section in sections:
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(30)
        section.right_margin = Mm(15)
        section.page_width = Mm(210)
        section.page_height = Mm(297)

        # Footer with page number (ГОСТ: на титульном не отображается)
        section.different_first_page_header_footer = True
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_ft.paragraph_format.space_before = Pt(0)
        p_ft.paragraph_format.space_after = Pt(0)
        r_ft = p_ft.add_run()
        r_ft.font.name = 'Times New Roman'
        r_ft.font.size = Pt(10)
        r_ft.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
        instr = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
        fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
        fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
        r_ft._r.append(fld1)
        r_ft._r.append(instr)
        r_ft._r.append(fld2)
        r_ft._r.append(fld3)

    # Styles
    styles = doc.styles

    # Normal font
    normal_style = styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # Colors
    STEAM_NAVY = RGBColor(0x1B, 0x28, 0x38)
    STEAM_BLUE = RGBColor(0x2D, 0x73, 0xA0)
    STEAM_CYAN = RGBColor(0x13, 0x58, 0x82)
    STEAM_GREEN = RGBColor(0x4C, 0x6B, 0x22)
    TEXT_DARK = RGBColor(0x1A, 0x1A, 0x1A)

    def set_cell_background(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def set_table_borders(table, color="B8B6B4", sz="4", val="single"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="none"/>'
            f'  <w:left w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_p(text="", style='Normal', bold=False, italic=False, color=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=4, indent=True):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if indent:
            p.paragraph_format.first_line_indent = Mm(12.5)
        else:
            p.paragraph_format.first_line_indent = Mm(0)

        if text:
            run = p.add_run(text)
            run.bold = bold
            run.italic = italic
            if color:
                run.font.color.rgb = color
            run.font.name = 'Times New Roman'
        return p

    def add_heading_1(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.first_line_indent = Mm(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(15)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = STEAM_NAVY
        return p

    def add_heading_2(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.first_line_indent = Mm(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(13)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = STEAM_BLUE
        return p

    def add_heading_3(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.first_line_indent = Mm(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(12)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = STEAM_CYAN
        return p

    def add_callout(text, prefix="Примечание: "):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Mm(165)
        set_cell_background(cell, "F0F4F8")
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

        # left border only
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="2D73A0"/>'
            f'  <w:top w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'  <w:bottom w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.first_line_indent = Mm(0)
        run_b = p.add_run(prefix)
        run_b.bold = True
        run_b.font.size = Pt(10.5)
        run_b.font.name = 'Times New Roman'
        run_b.font.color.rgb = STEAM_BLUE

        run_t = p.add_run(text)
        run_t.italic = True
        run_t.font.size = Pt(10.5)
        run_t.font.name = 'Times New Roman'
        run_t.font.color.rgb = TEXT_DARK

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_code_block(code_text, lang_title=""):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Mm(165)
        set_cell_background(cell, "F7F8FA")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:left w:val="single" w:sz="18" w:space="0" w:color="1B2838"/>'
            f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="D0D5DD"/>'
            f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="D0D5DD"/>'
            f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="D0D5DD"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.first_line_indent = Mm(0)

        if lang_title:
            run_title = p.add_run(lang_title + "\n")
            run_title.bold = True
            run_title.font.size = Pt(9)
            run_title.font.name = 'Segoe UI'
            run_title.font.color.rgb = STEAM_BLUE

        run_code = p.add_run(code_text)
        run_code.font.name = 'Consolas'
        run_code.font.size = Pt(8.5)
        run_code.font.color.rgb = RGBColor(0x20, 0x25, 0x30)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_figure(image_path, caption):
        if os.path.exists(image_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.paragraph_format.first_line_indent = Mm(0)
            p_img.paragraph_format.keep_with_next = True
            run = p_img.add_run()
            run.add_picture(image_path, width=Mm(155))

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(8)
            p_cap.paragraph_format.first_line_indent = Mm(0)
            run_cap = p_cap.add_run(caption)
            run_cap.font.name = 'Times New Roman'
            run_cap.font.size = Pt(10.5)
            run_cap.italic = True
            run_cap.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        else:
            add_p(f"[Изображение {image_path} не найдено]", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # -------------------------------------------------------------------------
    # 1. ТИТУЛЬНЫЙ ЛИСТ
    # -------------------------------------------------------------------------
    p_top = add_p("ДЕПАРТАМЕНТ ОБРАЗОВАНИЯ И НАУКИ ГОРОДА МОСКВЫ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=10, space_after=2, indent=False)
    p_top.runs[0].font.size = Pt(11)
    p_sub = add_p("ГБПОУ «КОЛЛЕДЖ СВЯЗИ И ИНФОРМАТИКИ»\nОтделение информационных технологий и программирования", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18, indent=False)
    p_sub.runs[0].font.size = Pt(10.5)

    add_p("Специальность: 09.02.07 «Информационные системы и программирование»", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_after=50, indent=False)

    p_doc = add_p("ОТЧЁТ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8, indent=False)
    p_doc.runs[0].font.size = Pt(22)
    p_doc.runs[0].font.color.rgb = STEAM_NAVY

    p_work = add_p("по практической разработке программного обеспечения", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12, indent=False)
    p_work.runs[0].font.size = Pt(13)

    p_theme = add_p("Тема: «Проектирование и разработка десктопного приложения цифровой дистрибуции игр и пополнения кошельков Steam \"KupiSteam\" на платформе WPF .NET»", 
                    bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=60, indent=False)
    p_theme.runs[0].font.size = Pt(14)
    p_theme.runs[0].font.color.rgb = STEAM_BLUE

    # Table for Author and Reviewer
    table_meta = doc.add_table(rows=2, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r in table_meta.rows:
        for c in r.cells:
            c.width = Mm(80)
            set_cell_margins(c, top=60, bottom=60, left=60, right=60)

    cell_l1 = table_meta.cell(0, 0)
    p = cell_l1.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Выполнил студент:")
    r.bold = True
    r.font.size = Pt(11)

    cell_r1 = table_meta.cell(0, 1)
    p = cell_r1.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Солодухин Владислав Артемьевич\nГруппа: ИСП-08")
    r.font.size = Pt(11)

    cell_l2 = table_meta.cell(1, 0)
    p = cell_l2.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Проверил преподаватель:")
    r.bold = True
    r.font.size = Pt(11)

    cell_r2 = table_meta.cell(1, 1)
    p = cell_r2.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Оценка: ___________ / Подпись: ______\n«____» ________________ 2026 г.")
    r.font.size = Pt(11)

    # City & Year at the bottom
    p_bot = add_p("\n\n\n\n\nМосква, 2026", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_before=40, space_after=0, indent=False)
    p_bot.runs[0].font.size = Pt(11)

    doc.add_page_break()

    # -------------------------------------------------------------------------
    # 2. СОДЕРЖАНИЕ
    # -------------------------------------------------------------------------
    add_heading_1("СОДЕРЖАНИЕ")

    toc_items = [
        ("ВВЕДЕНИЕ", "3"),
        ("1. ТЕХНИЧЕСКОЕ ЗАДАНИЕ НА РАЗРАБОТКУ ПО «KUPISTEAM»", "5"),
        ("   1.1. Основание разработки и назначение программного комплекса", "5"),
        ("   1.2. Требования к функциональным характеристикам", "5"),
        ("   1.3. Требования к надежности и сохранности данных", "7"),
        ("   1.4. Требования к аппаратному и программному обеспечению", "7"),
        ("2. АРХИТЕКТУРНОЕ ПРОЕКТИРОВАНИЕ И СХЕМЫ СИСТЕМЫ", "8"),
        ("   2.1. Трехуровневая архитектура приложения", "8"),
        ("   2.2. Диаграмма вариантов использования (Use Case Diagram)", "9"),
        ("   2.3. Диаграмма классов и моделирование данных", "10"),
        ("   2.4. Схема экранных форм и навигации интерфейса", "11"),
        ("   2.5. Алгоритмические схемы ключевых бизнес-процессов", "12"),
        ("3. ЭТАПЫ РАЗРАБОТКИ И ПРОГРАММНАЯ РЕАЛИЗАЦИЯ", "13"),
        ("   3.1. Этап 1. Предпроектный анализ и выбор технологического стека", "13"),
        ("   3.2. Этап 2. Проектирование визуальной дизайн-системы в стиле Steam", "14"),
        ("   3.3. Этап 3. Реализация модуля данных и файловой персистентности", "15"),
        ("   3.4. Этап 4. Реализация главного навигационного окна (MainWindow)", "17"),
        ("   3.5. Этап 5. Разработка модуля аутентификации и валидации (AuthWindow)", "18"),
        ("   3.6. Этап 6. Разработка каталога, баннерного слайдера и фильтрации", "20"),
        ("   3.7. Этап 7. Модуль корзины и дисконтных промокодов (CartPage)", "22"),
        ("   3.8. Этап 8. Модуль пополнения баланса Steam и расчет комиссии", "24"),
        ("   3.9. Этап 9. Личный кабинет, эмуляция банковских карт и история заказов", "25"),
        ("4. РУКОВОДСТВО ПОЛЬЗОВАТЕЛЯ И ДЕМОНСТРАЦИЯ РАБОТЫ (СКРИНШОТЫ)", "27"),
        ("   4.1. Авторизация и создание учетной записи", "27"),
        ("   4.2. Витрина цифрового магазина и динамический слайдер", "28"),
        ("   4.3. Страница прямого пополнения кошелька Steam", "29"),
        ("   4.4. Корзина покупок и активация промокодов", "30"),
        ("   4.5. Профиль пользователя и управление финансами", "31"),
        ("5. ТЕСТИРОВАНИЕ И ВЕРИФИКАЦИЯ ПРОГРАММНОГО ОБЕСПЕЧЕНИЯ", "32"),
        ("   5.1. Стратегия тестирования и сценарии проверки", "32"),
        ("   5.2. Сводная таблица результатов тестирования", "33"),
        ("ЗАКЛЮЧЕНИЕ", "35"),
        ("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", "36"),
    ]

    table_toc = doc.add_table(rows=len(toc_items), cols=2)
    table_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_toc, color="FFFFFF", val="none")

    for i, (item, page) in enumerate(toc_items):
        cell_title = table_toc.cell(i, 0)
        cell_page = table_toc.cell(i, 1)
        cell_title.width = Mm(150)
        cell_page.width = Mm(15)
        set_cell_margins(cell_title, top=40, bottom=40, left=0, right=0)
        set_cell_margins(cell_page, top=40, bottom=40, left=0, right=0)

        p_t = cell_title.paragraphs[0]
        p_t.paragraph_format.line_spacing = 1.15
        p_t.paragraph_format.space_after = Pt(2)
        p_t.paragraph_format.first_line_indent = Mm(0)
        run_t = p_t.add_run(item)
        run_t.font.name = 'Times New Roman'
        run_t.font.size = Pt(11)
        if item.isupper() and not item.startswith("   "):
            run_t.bold = True

        p_p = cell_page.paragraphs[0]
        p_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_p.paragraph_format.line_spacing = 1.15
        p_p.paragraph_format.space_after = Pt(2)
        p_p.paragraph_format.first_line_indent = Mm(0)
        run_p = p_p.add_run(page)
        run_p.font.name = 'Times New Roman'
        run_p.font.size = Pt(11)
        if item.isupper() and not item.startswith("   "):
            run_p.bold = True

    doc.add_page_break()

    # -------------------------------------------------------------------------
    # 3. ВВЕДЕНИЕ
    # -------------------------------------------------------------------------
    add_heading_1("ВВЕДЕНИЕ")

    add_heading_2("Актуальность темы")
    add_p(
        "Индустрия видеоигр является одним из наиболее динамично развивающихся секторов цифровой экономики. "
        "Ключевой мировой платформой дистрибуции компьютерных игр для персональных компьютеров на протяжении "
        "более двух десятилетий остается сервис Steam от компании Valve Corporation. "
        "Однако в современных реалиях российские пользователи столкнулись с беспрецедентными инфраструктурными "
        "и финансовыми ограничениями: прямая оплата с помощью российских банковских карт (МИР, Visa, MasterCard), "
        "а также популярных электронных кошельков внутри официального клиента Steam была заблокирована. "
        "Кроме того, ряд зарубежных издателей ограничил прямую продажу своих продуктов в российском регионе."
    )
    add_p(
        "В сложившейся ситуации возникла острая общественная и коммерческая потребность в надежных, быстрых и удобных "
        "отечественных программных инструментах, позволяющих российским геймерам легально приобретать ключи активации "
        "лицензионных игр, а также безопасно пополнять кошельки своих учетных записей Steam в рублях. "
        "Разработка десктопного программного комплекса «KupiSteam» направлена на решение указанной проблемы "
        "путем предоставления пользователю эргономичного интерфейса, выполненного в аутентичной стилистике официального клиента Steam, "
        "с интуитивно понятным каталогом, корзиной заказов, калькулятором пополнения баланса и личным кабинетом."
    )

    add_heading_2("Цели и задачи разработки")
    add_p(
        "Главной целью настоящей работы является проектирование, программная реализация и комплексное тестирование "
        "автономного десктопного приложения «KupiSteam» на базе платформы .NET Framework и технологии Windows Presentation Foundation (WPF), "
        "обеспечивающего полный цикл взаимодействия пользователя с каталогом цифровых товаров и сервисом пополнения игрового счета."
    )
    add_p("Для достижения поставленной цели в ходе выполнения работы решены следующие основные задачи:")
    
    tasks = [
        "Провести аналитический обзор существующих аналогов и сформулировать техническое задание на разработку ПО;",
        "Спроектировать архитектуру программного комплекса, диаграммы вариантов использования, классов и навигации экранных форм;",
        "Разработать модели данных (Game, CartItem, User, Order) и структуру локального файлового хранилища с гарантией целостности;",
        "Создать стилизованный графический интерфейс (XAML), воспроизводящий дизайн-код официального клиента Steam (цветовая палитра, типографика, градиенты, микроанимации);",
        "Реализовать бизнес-логику на языке C#: управление сессией пользователя, механизм авторизации и регистрации, поиск и фильтрация каталога, расчет корзины и промокодов, калькулятор комиссии;",
        "Разработать подсистему всплывающих Toast-уведомлений для повышения информативности и визуального отклика системы;",
        "Провести комплексное тестирование (функциональное, граничное, интеграционное) и составить техническую документацию."
    ]
    for t in tasks:
        p_t = add_p(f"•  {t}", indent=False, space_after=3)
        p_t.paragraph_format.left_indent = Mm(10)

    add_heading_2("Объект и предмет исследования")
    add_p("Объект исследования: процессы цифровой дистрибуции лицензионного игрового программного обеспечения и механизмы проведения электронных транзакций в десктопных средах.")
    add_p("Предмет исследования: методы и технологии создания пользовательского интерфейса и архитектуры приложений на стеке C# / WPF / XAML / .NET Framework.")

    add_callout(
        "Разработанный программный комплекс «KupiSteam» полностью автономен, не требует сторонних СУБД для развертывания, "
        "компилируется в среде Visual Studio с целевой платформой .NET Framework 4.7.2 и оптимизирован для работы на операционных системах Windows 10/11.",
        prefix="Ключевая особенность реализации: "
    )

    doc.add_page_break()

    # -------------------------------------------------------------------------
    # 4. РАЗДЕЛ 1. ТЕХНИЧЕСКОЕ ЗАДАНИЕ
    # -------------------------------------------------------------------------
    add_heading_1("1. ТЕХНИЧЕСКОЕ ЗАДАНИЕ НА РАЗРАБОТКУ ПО «KUPISTEAM»")

    add_heading_2("1.1. Основание разработки и назначение программного комплекса")
    add_p(
        "Настоящее техническое задание (ТЗ) регламентирует требования к функциональным возможностям, "
        "структуре данных, пользовательскому интерфейсу и условиям эксплуатации десктопного приложения «KupiSteam». "
        "Программный продукт предназначен для розничной реализации цифровых лицензионных ключей видеоигр "
        "и проведения операций пополнения рублевых кошельков платформы Steam через симуляцию платежного шлюза."
    )

    add_heading_2("1.2. Требования к функциональным характеристикам")
    add_p("Программный комплекс должен обеспечивать выполнение следующих ключевых функций, сгруппированных по подсистемам:")

    add_heading_3("1. Подсистема аутентификации и управления профилем:")
    p1 = [
        "Регистрация нового пользователя с проверкой уникальности логина, корректности адреса электронной почты и совпадения паролей;",
        "Вход зарегистрированного пользователя с валидацией учетных данных;",
        "Начисление приветственного бонуса в размере 500 рублей на внутренний баланс при успешной регистрации;",
        "Выход из учетной записи с мгновенным обновлением глобального состояния системы и шапки приложения;",
        "Возможность работы в режиме «Гость» с ограничением на подтверждение финансовых транзакций."
    ]
    for item in p1:
        p_i = add_p(f"—  {item}", indent=False, space_after=2)
        p_i.paragraph_format.left_indent = Mm(10)

    add_heading_3("2. Подсистема витрины и каталога игр:")
    p2 = [
        "Отображение промо-слайдера с популярными хитами (Cyberpunk 2077, Elden Ring, Baldur's Gate 3, Black Myth: Wukong, RDR 2), таймером автоматической смены (5 секунд) и ручным переключением;",
        "Фильтрация каталога по жанровым категориям («Все», «RPG», «Экшен», «Кооператив», «Инди», «Скидки»);",
        "Динамический («живой») поиск по названию игры по мере ввода символов в поле поиска;",
        "Отображение плашек скидок (-% от базовой цены), старой и актуальной цены;",
        "Быстрое добавление позиции в корзину непосредственно из карточки товара или промо-баннера."
    ]
    for item in p2:
        p_i = add_p(f"—  {item}", indent=False, space_after=2)
        p_i.paragraph_format.left_indent = Mm(10)

    add_heading_3("3. Подсистема корзины и применения промокодов:")
    p3 = [
        "Отображение перечня добавленных игр с постером, названием, жанром, ценой за единицу и итоговой стоимостью;",
        "Интерактивное управление количеством каждого товара (кнопки «+» и «-»), удаление отдельного товара, а также полная очистка корзины;",
        "Активация промокодов на скидку: промокод «STUDENT» (скидка 10%), промокод «KUPI20» (скидка 20%);",
        "Проверка достаточности средств на внутреннем кошельке пользователя перед завершением оформления заказа;",
        "Формирование записи заказа с генерацией уникальных лицензионных ключей и отправкой уведомления."
    ]
    for item in p3:
        p_i = add_p(f"—  {item}", indent=False, space_after=2)
        p_i.paragraph_format.left_indent = Mm(10)

    add_heading_3("4. Подсистема прямого пополнения кошелька Steam:")
    p4 = [
        "Ввод логина учетной записи Steam и желаемой суммы пополнения в диапазоне от 100 до 15 000 рублей;",
        "Кнопки быстрого выбора фиксированных номиналов (500, 1000, 2000, 5000 рублей);",
        "Автоматический расчет сервисного сбора (комиссии) в размере 5% и расчет итоговой суммы к списанию;",
        "Проверка корректности логина Steam, валидация лимитов и подтверждение перевода со списанием средств с кошелька."
    ]
    for item in p4:
        p_i = add_p(f"—  {item}", indent=False, space_after=2)
        p_i.paragraph_format.left_indent = Mm(10)

    add_heading_3("5. Подсистема личного кабинета и финансовых операций:")
    p5 = [
        "Отображение аватара пользователя (первая буква логина в фирменной градиентной рамке Steam), логина, почты и текущего баланса;",
        "Пополнение баланса аккаунта через симуляцию банковской карты с проверкой ввода 16-значного номера карты;",
        "Просмотр хронологической истории заказов и операций пополнения с указанием даты, времени, описания и суммы."
    ]
    for item in p5:
        p_i = add_p(f"—  {item}", indent=False, space_after=2)
        p_i.paragraph_format.left_indent = Mm(10)

    add_heading_2("1.3. Требования к надежности и сохранности данных")
    add_p(
        "Приложение должно обеспечивать устойчивость к некорректным действиям пользователя: "
        "блокировать ввод нечисловых символов в финансовые поля (PreviewTextInput), "
        "предотвращать создание пользователей с повторяющимися логинами или символом-разделителем «;», "
        "а также безопасно сохранять состояние базы данных пользователей (users.txt) и журнала транзакций (orders.txt) "
        "в кодировке UTF-8. В случае непредвиденных сбоев данные предыдущих сессий должны сохраняться."
    )

    add_heading_2("1.4. Требования к аппаратному и программному обеспечению")
    add_p("Таблица 1 определяет минимальные и рекомендуемые системные требования для запуска и функционирования ПО.")

    # Table of Requirements
    t_req = doc.add_table(rows=6, cols=3)
    t_req.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_req)

    req_data = [
        ("Параметр", "Минимальные требования", "Рекомендуемые требования"),
        ("Операционная система", "Microsoft Windows 7 SP1 / 8.1 / 10", "Microsoft Windows 10 / 11 (64-bit)"),
        ("Программная среда", "Microsoft .NET Framework 4.7.2", "Microsoft .NET Framework 4.8 / 4.8.1"),
        ("Процессор (CPU)", "1.6 GHz, 2 ядра (x86/x64)", "2.4 GHz+, 4 ядра (Intel Core i3 / AMD Ryzen 3+)"),
        ("Оперативная память (RAM)", "1 Гб свободной памяти", "2 Гб+ свободной памяти (приложение потребляет ~55 Мб)"),
        ("Разрешение экрана", "1150 x 650 пикселей", "1920 x 1080 пикселей (Full HD, 96/120 DPI)")
    ]

    for row_idx, row_values in enumerate(req_data):
        for col_idx, text in enumerate(row_values):
            cell = t_req.cell(row_idx, col_idx)
            if row_idx == 0:
                set_cell_background(cell, "1B2838")
                color = RGBColor(0xFF, 0xFF, 0xFF)
                bold = True
            else:
                set_cell_background(cell, "FFFFFF" if row_idx % 2 == 1 else "F7F8FA")
                color = TEXT_DARK
                bold = False
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.first_line_indent = Mm(0)
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            run.bold = bold
            run.font.color.rgb = color

    add_p("\nТаблица 1 — Системные требования программного комплекса «KupiSteam»", italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=10, indent=False)

    doc.add_page_break()

    # -------------------------------------------------------------------------
    # 5. РАЗДЕЛ 2. ПРОЕКТИРОВАНИЕ И СХЕМЫ
    # -------------------------------------------------------------------------
    add_heading_1("2. АРХИТЕКТУРНОЕ ПРОЕКТИРОВАНИЕ И СХЕМЫ СИСТЕМЫ")

    add_heading_2("2.1. Трехуровневая архитектура приложения")
    add_p(
        "Программный комплекс спроектирован на базе классической трехзвенной архитектуры "
        "(Presentation Layer — Business Logic Layer — Data Persistence Layer), дополненной принципами "
        "модульности и централизованного управления состоянием. Разделение обязанностей позволяет изолировать "
        "код графического интерфейса от алгоритмов расчета цен, применения скидок и записи в хранилище."
    )

    screens_dir = r"E:\ucheba\Kupisteam\Kupisteam\bin\Debug\Screenshots"
    add_figure(os.path.join(screens_dir, "diagram_architecture.png"), 
               "Рисунок 1 — Архитектурная схема программного комплекса «KupiSteam»")

    add_p(
        "На уровне представления (Presentation Layer) находятся окно MainWindow, содержащее постоянную верхнюю панель навигации "
        "и динамический контейнер Frame, модальное окно AuthWindow для входа и регистрации, а также набор автономных экранных страниц "
        "(CatalogPage, CartPage, TopUpPage, ProfilePage). Уровень бизнес-логики управляет процессами корзины, фильтрации, калькулятора "
        "и контроля баланса. Уровень персистентности отвечает за сериализацию и десериализацию сущностей в текстовые файлы."
    )

    add_heading_2("2.2. Диаграмма вариантов использования (Use Case Diagram)")
    add_p(
        "В системе выделено две основные роли акторов: «Гость» (неавторизованный посетитель) и «Авторизованный пользователь». "
        "Гость имеет право просматривать витрину, искать игры, изучать баннеры, знакомиться с условиями пополнения и наполнять корзину. "
        "Для совершения финансовых операций (оплата корзины, заказ пополнения Steam, зачисление средств) система требует прохождения авторизации."
    )

    add_figure(os.path.join(screens_dir, "diagram_usecase.png"), 
               "Рисунок 2 — Диаграмма вариантов использования (Use Case Diagram)")

    add_heading_2("2.3. Диаграмма классов и моделирование данных")
    add_p(
        "Информационная модель приложения строится вокруг четырех базовых сущностей: Game (игра каталога), "
        "CartItem (позиция в корзине покупок), User (учетная запись пользователя) и Order (транзакция в журнале заказов). "
        "Класс Data является статическим синглтоном, агрегирующим текущую сессию пользователя, корзину и методы ввода-вывода."
    )

    add_figure(os.path.join(screens_dir, "diagram_datamodel.png"), 
               "Рисунок 3 — Диаграмма классов и сущностей модели данных")

    add_p(
        "Связь между сущностями Game и CartItem выражается отношением агрегации «1 ко многим»: объект CartItem инкапсулирует ссылку "
        "на объект Game и целое число Count. Сущность User связана с историей транзакций Order через строковый идентификатор Login."
    )

    add_heading_2("2.4. Схема экранных форм и навигации интерфейса")
    add_p(
        "Навигационная модель построена по принципу Single Window Application с внедрением компонента Frame. "
        "Верхний колонтитул (Header) окна MainWindow остается персистентным на протяжении всего жизненного цикла программы, "
        "отображая актуальный баланс, индикатор количества позиций в корзине и состояние сессии. Навигация между страницами "
        "осуществляется бесшовно без создания лишних окон операционной системы."
    )

    add_figure(os.path.join(screens_dir, "diagram_navigation.png"), 
               "Рисунок 4 — Схема экранных форм и переходов пользовательского интерфейса")

    add_heading_2("2.5. Алгоритмические схемы ключевых бизнес-процессов")
    add_p(
        "Наиболее критически важным алгоритмом является процесс оформления покупки товаров из корзины. "
        "Алгоритм включает последовательность проверок: наличие активной авторизации (при ее отсутствии открывается AuthWindow), "
        "валидация достаточности баланса кошелька относительно итоговой стоимости с учетом примененного промокода, "
        "атомарное списание средств, запись транзакции и очистка корзины."
    )

    add_figure(os.path.join(screens_dir, "diagram_algorithm_checkout.png"), 
               "Рисунок 5 — Блок-схема алгоритма оформления заказа и проверки баланса")

    doc.add_page_break()

    # -------------------------------------------------------------------------
    # 6. РАЗДЕЛ 3. ЭТАПЫ РАЗРАБОТКИ И ПРОГРАММНАЯ РЕАЛИЗАЦИЯ
    # -------------------------------------------------------------------------
    add_heading_1("3. ЭТАПЫ РАЗРАБОТКИ И ПРОГРАММНАЯ РЕАЛИЗАЦИЯ")

    add_heading_2("3.1. Этап 1. Предпроектный анализ и выбор технологического стека")
    add_p(
        "На начальном этапе разработки был выполнен сравнительный анализ инструментов создания настольных приложений в среде Windows: "
        "Windows Forms, WPF (Windows Presentation Foundation), Electron и MAUI. Выбор технологии WPF обусловлен следующими факторами:"
    )
    p_reasons = [
        "Аппаратное ускорение графики через DirectX, обеспечивающее высокую плавность отрисовки (60 FPS);",
        "Декларативное описание интерфейса на языке XAML, позволяющее отделить визуальную разметку от алгоритмического кода;",
        "Мощная система стилей, шаблонов элементов управления (ControlTemplate) и триггеров;",
        "Поддержка векторной графики, растровых эффектов (DropShadowEffect) и покадровых анимаций (DoubleAnimation);",
        "Высокая производительность среды выполнения .NET Framework и отсутствие оверхеда Chromium-движка (в отличие от Electron)."
    ]
    for r in p_reasons:
        p_i = add_p(f"1.  {r}", indent=False, space_after=2)
        p_i.paragraph_format.left_indent = Mm(10)

    add_heading_2("3.2. Этап 2. Проектирование визуальной дизайн-системы в стиле Steam")
    add_p(
        "Для достижения максимальной аутентичности интерфейс приложения спроектирован в фирменной темной палитре Steam. "
        "Все цветовые константы, градиенты, стили кнопок, полей ввода и карточек вынесены в файл глобальных ресурсов App.xaml."
    )

    add_callout(
        "Базовая палитра включает: глубокий сине-черный фон #1B2838, темную шапку #171A21, карточки #16202D, "
        "фирменный голубой акцент #66C0F4, ярко-зеленый оттенок скидок #A4D007, цены со скидкой #BEEE11 и плашки #4C6B22.",
        prefix="Дизайн-токены Steam: "
    )

    add_p("Ниже представлен фрагмент определения глобальных кистей и стиля фирменной зеленой кнопки оформления заказа:")

    code_styles = (
        '<!-- App.xaml: Палитра цветов и стиль кнопки в стиле Steam -->\n'
        '<SolidColorBrush x:Key="BgBrush" Color="#1B2838"/>\n'
        '<SolidColorBrush x:Key="HeaderBrush" Color="#171A21"/>\n'
        '<SolidColorBrush x:Key="LinkBrush" Color="#66C0F4"/>\n'
        '<SolidColorBrush x:Key="PriceGreenBrush" Color="#BEEE11"/>\n'
        '<SolidColorBrush x:Key="DiscountBgBrush" Color="#4C6B22"/>\n\n'
        '<!-- Стиль кнопки действия AccentButton (градиент и hover-эффект) -->\n'
        '<Style x:Key="AccentButton" TargetType="Button">\n'
        '    <Setter Property="Foreground" Value="#D2EFA9"/>\n'
        '    <Setter Property="FontFamily" Value="Segoe UI"/>\n'
        '    <Setter Property="FontSize" Value="14"/>\n'
        '    <Setter Property="Padding" Value="16,8"/>\n'
        '    <Setter Property="Cursor" Value="Hand"/>\n'
        '    <Setter Property="Background">\n'
        '        <Setter.Value>\n'
        '            <LinearGradientBrush StartPoint="0,0" EndPoint="1,0">\n'
        '                <GradientStop Color="#75B022" Offset="0"/>\n'
        '                <GradientStop Color="#588A1B" Offset="1"/>\n'
        '            </LinearGradientBrush>\n'
        '        </Setter.Value>\n'
        '    </Setter>\n'
        '    <Setter Property="Template">\n'
        '        <Setter.Value>\n'
        '            <ControlTemplate TargetType="Button">\n'
        '                <Border x:Name="bd" Background="{TemplateBinding Background}" CornerRadius="2" Padding="{TemplateBinding Padding}">\n'
        '                    <ContentPresenter HorizontalAlignment="Center" VerticalAlignment="Center"/>\n'
        '                </Border>\n'
        '                <ControlTemplate.Triggers>\n'
        '                    <Trigger Property="IsMouseOver" Value="True">\n'
        '                        <Setter TargetName="bd" Property="Background">\n'
        '                            <Setter.Value>\n'
        '                                <LinearGradientBrush StartPoint="0,0" EndPoint="1,0">\n'
        '                                    <GradientStop Color="#8ED229" Offset="0"/>\n'
        '                                    <GradientStop Color="#6DA821" Offset="1"/>\n'
        '                                </LinearGradientBrush>\n'
        '                            </Setter.Value>\n'
        '                        </Setter>\n'
        '                    </Trigger>\n'
        '                </ControlTemplate.Triggers>\n'
        '            </ControlTemplate>\n'
        '        </Setter.Value>\n'
        '    </Setter>\n'
        '</Style>'
    )
    add_code_block(code_styles, "Листинг 1 — Стилизация интерфейса в стиле Steam (App.xaml)")

    add_heading_2("3.3. Этап 3. Реализация модуля данных и файловой персистентности")
    add_p(
        "Класс Data.cs выступает единой точкой доступа к состоянию приложения. "
        "Хранение учетных записей реализовано в файле users.txt, где каждая строка имеет структурированный формат: "
        "«логин;почта;пароль;баланс». Журнал заказов хранится в orders.txt в формате «логин;дата;описание;сумма»."
    )

    code_data = (
        '// Data.cs: Модели и механизмы сохранения состояния\n'
        'public class Game\n'
        '{\n'
        '    public string Name { get; set; }\n'
        '    public string Genre { get; set; }\n'
        '    public int Price { get; set; }\n'
        '    public int OldPrice { get; set; }\n'
        '    public string Image { get; set; }\n\n'
        '    public Game(string name, string genre, int price, int oldPrice, int steamId)\n'
        '    {\n'
        '        Name = name; Genre = genre; Price = price; OldPrice = oldPrice;\n'
        '        Image = "pack://application:,,,/Images/game_" + steamId + ".jpg";\n'
        '    }\n'
        '    public bool HasDiscount => OldPrice > Price;\n'
        '    public string DiscountText => HasDiscount ? "-" + (100 - Price * 100 / OldPrice) + "%" : "";\n'
        '    public string PriceText => Price + " ₽";\n'
        '}\n\n'
        'public static class Data\n'
        '{\n'
        '    public static User CurrentUser = null;\n'
        '    public static List<CartItem> Cart = new List<CartItem>();\n'
        '    static string usersFile = AppDomain.CurrentDomain.BaseDirectory + "users.txt";\n'
        '    static string ordersFile = AppDomain.CurrentDomain.BaseDirectory + "orders.txt";\n\n'
        '    public static List<User> LoadUsers()\n'
        '    {\n'
        '        List<User> users = new List<User>();\n'
        '        if (!File.Exists(usersFile)) return users;\n'
        '        foreach (string line in File.ReadAllLines(usersFile, Encoding.UTF8))\n'
        '        {\n'
        '            string[] p = line.Split(\';\');\n'
        '            if (p.Length >= 4)\n'
        '                users.Add(new User { Login = p[0], Email = p[1], Password = p[2], Balance = int.Parse(p[3]) });\n'
        '        }\n'
        '        return users;\n'
        '    }\n'
        '    public static void SaveCurrentUser()\n'
        '    {\n'
        '        List<User> users = LoadUsers();\n'
        '        for (int i = 0; i < users.Count; i++)\n'
        '            if (users[i].Login == CurrentUser.Login) users[i] = CurrentUser;\n'
        '        SaveUsers(users);\n'
        '    }\n'
        '}'
    )
    add_code_block(code_data, "Листинг 2 — Структуры данных и методы персистентности (Data.cs)")

    add_heading_2("3.4. Этап 4. Реализация главного навигационного окна (MainWindow)")
    add_p(
        "Класс MainWindow обеспечивает глобальное состояние шапки приложения, переключение вкладок меню, "
        "вызов модального окна входа и отображение всплывающих уведомлений Toast. "
        "Анимация появления и затухания уведомления реализована с помощью DoubleAnimationUsingKeyFrames."
    )

    code_main = (
        '// MainWindow.xaml.cs: Управление шапкой и Toast-уведомлением\n'
        'public void UpdateHeader()\n'
        '{\n'
        '    int count = Data.CartCount();\n'
        '    CartCountText.Text = count.ToString();\n'
        '    CartBadge.Visibility = count > 0 ? Visibility.Visible : Visibility.Collapsed;\n\n'
        '    if (Data.CurrentUser != null)\n'
        '    {\n'
        '        LoginButton.Visibility = Visibility.Collapsed;\n'
        '        ProfileButton.Visibility = Visibility.Visible;\n'
        '        BalancePanel.Visibility = Visibility.Visible;\n'
        '        ProfileName.Text = Data.CurrentUser.Login;\n'
        '        HeaderAvatar.Text = Data.CurrentUser.Login.Substring(0, 1).ToUpper();\n'
        '        BalanceText.Text = Data.CurrentUser.Balance + " ₽";\n'
        '    }\n'
        '    else\n'
        '    {\n'
        '        LoginButton.Visibility = Visibility.Visible;\n'
        '        ProfileButton.Visibility = Visibility.Collapsed;\n'
        '        BalancePanel.Visibility = Visibility.Collapsed;\n'
        '    }\n'
        '}\n\n'
        'public void ShowToast(string text)\n'
        '{\n'
        '    ToastText.Text = text;\n'
        '    DoubleAnimationUsingKeyFrames anim = new DoubleAnimationUsingKeyFrames();\n'
        '    anim.KeyFrames.Add(new LinearDoubleKeyFrame(1, KeyTime.FromTimeSpan(TimeSpan.FromSeconds(0.2))));\n'
        '    anim.KeyFrames.Add(new LinearDoubleKeyFrame(1, KeyTime.FromTimeSpan(TimeSpan.FromSeconds(2.2))));\n'
        '    anim.KeyFrames.Add(new LinearDoubleKeyFrame(0, KeyTime.FromTimeSpan(TimeSpan.FromSeconds(2.6))));\n'
        '    Toast.BeginAnimation(OpacityProperty, anim);\n'
        '}'
    )
    add_code_block(code_main, "Листинг 3 — Логика обновления интерфейса и анимации уведомлений (MainWindow.xaml.cs)")

    add_heading_2("3.5. Этап 5. Разработка модуля аутентификации и валидации (AuthWindow)")
    add_p(
        "Окно AuthWindow спроектировано без стандартной рамки операционной системы (WindowStyle=\"None\", AllowsTransparency=\"True\"), "
        "поддерживает перетаскивание курсором мыши (DragMove) и переключение вкладок «Вход» и «Создать аккаунт». "
        "При регистрации выполняется строгая верификация: запрет служебных символов, валидация почты по наличию знаков '@' и '.', "
        "проверка совпадения паролей и начисление стартового депозита 500 рублей."
    )

    code_auth = (
        '// AuthWindow.xaml.cs: Алгоритм регистрации с валидацией\n'
        'private void DoRegister_Click(object sender, RoutedEventArgs e)\n'
        '{\n'
        '    string login = RegLogin.Text.Trim();\n'
        '    string email = RegEmail.Text.Trim();\n'
        '    string pass = RegPassword.Password;\n'
        '    string pass2 = RegPassword2.Password;\n\n'
        '    if (login == "" || email == "" || pass == "") { RegError.Text = "Заполните все поля"; return; }\n'
        '    if (login.Length < 3 || login.Contains(";")) { RegError.Text = "Логин от 3 символов (без «;»)"; return; }\n'
        '    if (!email.Contains("@") || !email.Contains(".")) { RegError.Text = "Введите корректную почту"; return; }\n'
        '    if (pass.Length < 4 || pass.Contains(";")) { RegError.Text = "Пароль не короче 4 символов"; return; }\n'
        '    if (pass != pass2) { RegError.Text = "Пароли не совпадают"; return; }\n'
        '    if (Data.FindUser(login) != null) { RegError.Text = "Такой логин уже занят"; return; }\n\n'
        '    User user = new User { Login = login, Email = email, Password = pass, Balance = 500 };\n'
        '    Data.AddUser(user);\n'
        '    Data.CurrentUser = user;\n'
        '    MessageBox.Show("Аккаунт создан! Начислен бонус 500 ₽.", "KupiSteam", MessageBoxButton.OK, MessageBoxImage.Information);\n'
        '    Close();\n'
        '}'
    )
    add_code_block(code_auth, "Листинг 4 — Алгоритм верификации и регистрации пользователя (AuthWindow.xaml.cs)")

    add_heading_2("3.6. Этап 6. Разработка каталога, баннерного слайдера и фильтрации")
    add_p(
        "На странице CatalogPage реализован карусельный промо-слайдер новинок и скидок с использованием таймера DispatcherTimer, "
        "динамические индикаторы слайдов (Dots) и фильтрация каталога игр. Метод ApplyFilter() осуществляет комбинированную "
        "выборку по выбранному жанру и подстроке поискового запроса."
    )

    code_cat = (
        '// CatalogPage.xaml.cs: Фильтрация игр и слайдер баннеров\n'
        'void ApplyFilter()\n'
        '{\n'
        '    string search = SearchBox.Text.Trim().ToLower();\n'
        '    List<Game> result = new List<Game>();\n\n'
        '    foreach (Game g in Data.Games)\n'
        '    {\n'
        '        bool genreOk = currentGenre == "Все"\n'
        '                       || g.Genre == currentGenre\n'
        '                       || (currentGenre == "Скидки" && g.HasDiscount);\n'
        '        bool searchOk = g.Name.ToLower().Contains(search);\n\n'
        '        if (genreOk && searchOk)\n'
        '            result.Add(g);\n'
        '    }\n'
        '    GamesList.ItemsSource = result;\n'
        '    FoundText.Text = "найдено игр: " + result.Count;\n'
        '    EmptyText.Visibility = result.Count == 0 ? Visibility.Visible : Visibility.Collapsed;\n'
        '}'
    )
    add_code_block(code_cat, "Листинг 5 — Логика фильтрации и поиска по каталогу игр (CatalogPage.xaml.cs)")

    add_heading_2("3.7. Этап 7. Модуль корзины и дисконтных промокодов (CartPage)")
    add_p(
        "Страница CartPage отображает список выбранных позиций, обеспечивает модификацию счетчиков количества, "
        "проверяет валидность дисконтных купонов и выполняет атомарное списание средств при оформлении заказа."
    )

    code_cart = (
        '// CartPage.xaml.cs: Промокоды и расчет итоговой стоимости\n'
        'private void Promo_Click(object sender, RoutedEventArgs e)\n'
        '{\n'
        '    string code = PromoBox.Text.Trim().ToUpper();\n'
        '    if (code == "STUDENT") promoPercent = 10;\n'
        '    else if (code == "KUPI20") promoPercent = 20;\n'
        '    else promoPercent = 0;\n\n'
        '    if (promoPercent > 0) {\n'
        '        PromoInfo.Text = "Промокод применён: -" + promoPercent + "%";\n'
        '        PromoInfo.Foreground = (Brush)FindResource("GreenBrush");\n'
        '    } else {\n'
        '        PromoInfo.Text = "Такого промокода нет";\n'
        '        PromoInfo.Foreground = (Brush)FindResource("ErrorBrush");\n'
        '    }\n'
        '    Refresh();\n'
        '}'
    )
    add_code_block(code_cart, "Листинг 6 — Активация промокодов и расчет скидки (CartPage.xaml.cs)")

    add_heading_2("3.8. Этап 8. Модуль пополнения баланса Steam и расчет комиссии")
    add_p(
        "Страница TopUpPage обеспечивает конвертацию рублей в баланс кошелька Steam с прозрачным расчетом "
        "комиссионного вознаграждения сервиса (5%). Метод Calculate() динамически обновляет суммы к зачислению "
        "и списанию при каждом изменении поля ввода."
    )

    code_topup = (
        '// TopUpPage.xaml.cs: Расчет суммы пополнения и комиссии сервиса\n'
        'const int MinSum = 100, MaxSum = 15000, FeePercent = 5;\n\n'
        'void Calculate()\n'
        '{\n'
        '    int amount;\n'
        '    int.TryParse(AmountBox.Text, out amount);\n'
        '    int fee = amount * FeePercent / 100;\n'
        '    int total = amount + fee;\n\n'
        '    ReceiveText.Text = amount + " ₽";\n'
        '    FeeText.Text = fee + " ₽";\n'
        '    TotalText.Text = total + " ₽";\n'
        '    UserBalanceText.Text = Data.CurrentUser != null ? Data.CurrentUser.Balance + " ₽" : "войдите в аккаунт";\n'
        '}'
    )
    add_code_block(code_topup, "Листинг 7 — Калькулятор сервисного сбора (TopUpPage.xaml.cs)")

    add_heading_2("3.9. Этап 9. Личный кабинет, эмуляция банковских карт и история заказов")
    add_p(
        "В модуле ProfilePage реализован механизм пополнения внутреннего кошелька пользователя. "
        "Проверяется формат номера банковской карты (ровно 16 цифр), допустимый диапазон суммы (от 10 до 100 000 руб.), "
        "после чего выполняется зачисление и мгновенное обновление хронологической истории операций."
    )

    doc.add_page_break()

    # -------------------------------------------------------------------------
    # 7. РАЗДЕЛ 4. РУКОВОДСТВО ПОЛЬЗОВАТЕЛЯ И СКРИНШОТЫ
    # -------------------------------------------------------------------------
    add_heading_1("4. РУКОВОДСТВО ПОЛЬЗОВАТЕЛЯ И ДЕМОНСТРАЦИЯ РАБОТЫ (СКРИНШОТЫ)")

    add_heading_2("4.1. Авторизация и создание учетной записи")
    add_p(
        "При запуске приложения или нажатии на кнопку «войти» в правом верхнем углу открывается "
        "стилизованное модальное диалоговое окно AuthWindow. Окно имеет две вкладки: «ВХОД» и «СОЗДАТЬ АККАУНТ». "
        "Вкладки переключаются кликом мыши, подсвечивая активное состояние."
    )

    add_figure(os.path.join(screens_dir, "01_auth_login.png"), 
               "Рисунок 6 — Модальное окно авторизации пользователя (Вкладка «Вход»)")

    add_p(
        "При переходе на вкладку регистрации пользователь указывает уникальный логин, рабочий адрес электронной почты "
        "и дважды вводит пароль для подтверждения. При успешной регистрации на баланс аккаунта автоматически зачисляется "
        "приветственный бонус в размере 500 рублей, о чем информирует всплывающее системное сообщение."
    )

    add_figure(os.path.join(screens_dir, "02_auth_register.png"), 
               "Рисунок 7 — Форма создания новой учетной записи с приветственным бонусом")

    add_heading_2("4.2. Витрина цифрового магазина и динамический слайдер")
    add_p(
        "Главный экран приложения (CatalogPage) встречает пользователя интерактивным широкоформатным промо-слайдером. "
        "Слайдер содержит скриншот игры, бейдж («ХИТ ПРОДАЖ», «СКИДКА -33%», «ИГРА ГОДА»), краткое описание, цену и зеленую кнопку «КУПИТЬ». "
        "Ниже расположены кнопки фильтрации по жанрам, строка поиска и сетка карточек игр каталога."
    )

    add_figure(os.path.join(screens_dir, "03_catalog_page.png"), 
               "Рисунок 8 — Главный экран каталога игр «KupiSteam» со слайдером и фильтрами")

    add_heading_2("4.3. Страница прямого пополнения кошелька Steam")
    add_p(
        "Вкладка «ПОПОЛНЕНИЕ STEAM» открывает эргономичный платежный модуль. Пользователю предлагается ввести логин аккаунта Steam "
        "и указать сумму пополнения. Для удобства предусмотрены кнопки быстрого выбора номинала (500, 1000, 2000, 5000 ₽). "
        "В правой колонке интерактивно отображается подробная сводка платежа с расчетом комиссии 5% и сверкой баланса пользователя."
    )

    add_figure(os.path.join(screens_dir, "04_topup_page.png"), 
               "Рисунок 9 — Экран пополнения кошелька Steam с автоматическим расчетом комиссии")

    add_heading_2("4.4. Корзина покупок и активация промокодов")
    add_p(
        "В корзине (CartPage) пользователь может гибко управлять составом своего заказа: увеличивать или уменьшать количество копий, "
        "удалять ненужные позиции. В правой панели предусмотрено поле ввода промокодов на скидку: "
        "при вводе «STUDENT» применяется скидка 10%, при вводе «KUPI20» — скидка 20%. Итоговая сумма пересчитывается мгновенно."
    )

    add_figure(os.path.join(screens_dir, "05_cart_page.png"), 
               "Рисунок 10 — Корзина товаров с примененным промокодом на скидку")

    add_heading_2("4.5. Профиль пользователя и управление финансами")
    add_p(
        "Экран профиля (ProfilePage) объединяет паспортные данные учетной записи, форму пополнения баланса с банковской карты "
        "и подробную хронологическую таблицу всех завершенных заказов и транзакций с указанием точной даты и времени."
    )

    add_figure(os.path.join(screens_dir, "06_profile_page.png"), 
               "Рисунок 11 — Личный кабинет пользователя с историей покупок и формой пополнения картой")

    doc.add_page_break()

    # -------------------------------------------------------------------------
    # 8. РАЗДЕЛ 5. ТЕСТИРОВАНИЕ И ВЕРИФИКАЦИЯ
    # -------------------------------------------------------------------------
    add_heading_1("5. ТЕСТИРОВАНИЕ И ВЕРИФИКАЦИЯ ПРОГРАММНОГО ОБЕСПЕЧЕНИЯ")

    add_heading_2("5.1. Стратегия тестирования и сценарии проверки")
    add_p(
        "Для проверки корректности работы программного комплекса «KupiSteam» было проведено комплексное тестирование, "
        "включающее методы черного ящика (Black-Box Testing), проверку граничных условий (Boundary Value Analysis), "
        "а также стресс-тестирование пользовательского ввода. В Таблице 2 сведены ключевые тест-кейсы."
    )

    # Test cases table
    t_test = doc.add_table(rows=11, cols=5)
    t_test.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_test)

    headers = ["№", "Тестируемый модуль", "Входные данные / Действие", "Ожидаемый результат", "Статус"]
    widths = [Mm(8), Mm(32), Mm(45), Mm(60), Mm(20)]

    for col_idx, h in enumerate(headers):
        cell = t_test.cell(0, col_idx)
        cell.width = widths[col_idx]
        set_cell_background(cell, "1B2838")
        set_cell_margins(cell, top=80, bottom=80, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Mm(0)
        run = p.add_run(h)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    tests_data = [
        ("1", "Регистрация", "Логин: 'al', пароли: '123'", "Ошибка: логин от 3 симв., пароль от 4", "Пройден"),
        ("2", "Регистрация", "Корректные данные, новый логин", "Успешная регистрация + бонус 500 ₽", "Пройден"),
        ("3", "Регистрация", "Логин содержит разделитель ';'", "Ошибка валидации служебного символа", "Пройден"),
        ("4", "Каталог", "Ввод 'cyber' в строку поиска", "Отображение Cyberpunk 2077, счетчик: 1", "Пройден"),
        ("5", "Каталог", "Выбор категории 'Инди'", "Фильтрация: Hades, Terraria, Stardew Valley", "Пройден"),
        ("6", "Корзина", "Ввод промокода 'STUDENT'", "Применение скидки 10%, пересчет итога", "Пройден"),
        ("7", "Корзина", "Оформление при нехватке средств", "Блокировка, подсказка пополнить баланс", "Пройден"),
        ("8", "Корзина", "Оформление при достаточном балансе", "Списание средств, заказ в истории, MessageBox", "Пройден"),
        ("9", "Пополнение Steam", "Ввод суммы 50 ₽ (меньше минимума)", "Ошибка: сумма должна быть от 100 до 15 000 ₽", "Пройден"),
        ("10", "Профиль", "Ввод 15 цифр карты вместо 16", "Ошибка: номер карты должен содержать 16 цифр", "Пройден")
    ]

    for row_idx, data in enumerate(tests_data, start=1):
        for col_idx, text in enumerate(data):
            cell = t_test.cell(row_idx, col_idx)
            cell.width = widths[col_idx]
            bg_color = "FFFFFF" if row_idx % 2 == 1 else "F7F8FA"
            if col_idx == 4:
                bg_color = "EAF5D6" # light green for passed
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 4] else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Mm(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(8.5)
            if col_idx == 4:
                run.bold = True
                run.font.color.rgb = STEAM_GREEN

    add_p("\nТаблица 2 — Результаты функционального и регрессионного тестирования", italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=10, indent=False)

    doc.add_page_break()

    # -------------------------------------------------------------------------
    # 9. ЗАКЛЮЧЕНИЕ
    # -------------------------------------------------------------------------
    add_heading_1("ЗАКЛЮЧЕНИЕ")
    add_p(
        "В результате выполнения практической работы было успешно спроектировано, разработано и всесторонне протестировано "
        "десктопное приложение цифровой дистрибуции видеоигр и сервисного пополнения баланса «KupiSteam». "
        "Все цели и задачи, поставленные в техническом задании, были достигнуты в полном объеме."
    )
    add_p("В процессе разработки были получены следующие практические результаты:")
    
    results = [
        "Спроектирована и программно реализована гибкая трехзвенная архитектура программного комплекса на базе технологии WPF и языка C#;",
        "Воссоздан фирменный визуальный дизайн в стилистике клиента Steam с плавной анимацией Toast-оповещений и отзывчивым управлением;",
        "Реализован полнофункциональный каталог видеоигр с динамическим баннерным слайдером, многофакторной фильтрацией и живым поиском;",
        "Создан модуль корзины с поддержкой процентных скидочных промокодов, калькулятора сумм и механизмов защиты от переполнения;",
        "Разработан надежный калькулятор пополнения баланса Steam с расчетом сервисного сбора (5%) и верификацией допустимых номиналов;",
        "Обеспечена надежная система хранения учетных записей и истории заказов с защитой от повреждения структуры файлов."
    ]
    for r in results:
        p_i = add_p(f"✓  {r}", indent=False, space_after=2)
        p_i.paragraph_format.left_indent = Mm(10)

    add_heading_2("Перспективы дальнейшего развития проекта")
    add_p(
        "Разработанное программное обеспечение обладает высоким потенциалом для масштабирования. В качестве направлений дальнейшей модернизации выделяются: "
        "переход от локальных файлов к реляционной базе данных (PostgreSQL / SQLite через Entity Framework Core), "
        "интеграция с официальным Steam Web API для автоматической валидации существующих аккаунтов и загрузки актуальных аватаров, "
        "подключение реального шлюза онлайн-эквайринга (ЮKassa / СБП), а также внедрение системы отзывов и оценок игроков."
    )

    doc.add_page_break()

    # -------------------------------------------------------------------------
    # 10. СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ
    # -------------------------------------------------------------------------
    add_heading_1("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")

    sources = [
        "ГОСТ 19.201-78. Единая система программной документации. Техническое задание. Требования к содержанию и оформлению. — М.: Стандартинформ, 2010.",
        "ГОСТ 7.32-2017. Отчет о научно-исследовательской работе. Структура и правила оформления. — М.: Стандартинформ, 2018.",
        "Троелсен, Э. Язык программирования C# 7 и платформы .NET и .NET Core / Э. Троелсен, Ф. Джепикс. — 8-е изд. — М.: Диалектика-Вильямс, 2018. — 1328 с.",
        "Макдональд, М. WPF: Windows Presentation Foundation в .NET 4.5 с примерами на C# / М. Макдональд. — 4-е изд. — СПб.: БХВ-Петербург, 2017. — 1024 с.",
        "Рихтер, Дж. CLR via C#. Программирование на платформе Microsoft .NET Framework 4.5 на языке C# / Дж. Рихтер. — 4-е изд. — СПб.: Питер, 2019. — 896 с.",
        "Официальная документация Microsoft Learn: Документация по Windows Presentation Foundation (WPF) [Электронный ресурс]. — Режим доступа: https://learn.microsoft.com/ru-ru/dotnet/desktop/wpf/ (дата обращения: 24.10.2026).",
        "Официальная документация Microsoft Learn: Обзор архитектуры XAML и разметки [Электронный ресурс]. — Режим доступа: https://learn.microsoft.com/ru-ru/dotnet/desktop/wpf/xaml/ (дата обращения: 24.10.2026).",
        "Valve Corporation. Steamworks Documentation and Design Guidelines [Электронный ресурс]. — Режим доступа: https://partner.steamgames.com/doc/home (дата обращения: 24.10.2026)."
    ]

    for idx, s in enumerate(sources, start=1):
        p_s = add_p(f"{idx}.  {s}", indent=False, space_after=4)
        p_s.paragraph_format.left_indent = Mm(10)
        p_s.paragraph_format.first_line_indent = Mm(-10)

    # Output file
    output_docx = r"E:\ucheba\Kupisteam\Отчет_KupiSteam_Солодухин_ИСП-08.docx"
    doc.save(output_docx)
    print("SUCCESS: Report saved to", output_docx)

if __name__ == '__main__':
    create_report()
